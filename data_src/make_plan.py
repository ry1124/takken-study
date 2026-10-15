# -*- coding: utf-8 -*-
"""宅建 2027年合格プラン(厳しめ版)の毎日の計画を作る。
出力: data_src/plan.json(アプリ用) と 宅建\\宅建_毎日の学習計画.xlsx(人が見る用)。
方針: 毎日30分は宅建業法の一問一答 / 章を読んだ日に同じ範囲の過去問 / 日曜は週テスト(完全休養は月1回) /
      月末に関門(届かなければ次に進まず翌週を復習) / 5月から毎週土曜に年度別過去問の通し演習。"""
import datetime as dt
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT_JSON = HERE / "plan.json"
OUT_XLSX = HERE.parent.parent / "宅建" / "宅建_毎日の学習計画.xlsx"
D = dt.date
WD = "月火水木金土日"
START, EXAM = D(2026, 10, 16), D(2027, 10, 17)
HOLIDAYS = {D(2026, 12, 31), D(2027, 1, 1)}   # 年末年始は休み

IPPON = "①宅建業法の一問一答（30分）"
FIX = "③×の問題を解き直し、なぜ間違えたかを1行メモ（30分）"

GYOHO = [("宅建業の定義（宅地・建物・取引・業）", 1), ("免許の種類・免許換え・届出", 2), ("免許の欠格事由", 2),
         ("事務所・案内所・標識・帳簿", 1), ("宅建士の登録・宅建士証", 2), ("登録の移転・変更の登録", 1),
         ("営業保証金", 2), ("保証協会（弁済業務保証金）", 2), ("媒介契約（一般・専任・専属専任）", 2),
         ("広告の規制・取引態様の明示", 1), ("重要事項説明（35条）基本", 2), ("重要事項説明（35条）区分所有・貸借", 2),
         ("37条書面（契約書）", 2), ("8種制限：クーリング・オフ", 2), ("8種制限：手付金等の保全措置", 2),
         ("8種制限：損害賠償額の予定・手付・他人物売買", 1), ("8種制限：担保責任の特約・割賦販売", 1),
         ("報酬額の制限（売買）", 2), ("報酬額の制限（貸借）", 1), ("その他の業務上の規制（禁止行為）", 1),
         ("監督処分・罰則", 2), ("住宅瑕疵担保履行法", 1)]
HOREI = [("都市計画法：区域区分・都市計画の種類", 2), ("都市計画法：地域地区・地区計画", 1),
         ("都市計画法：都市計画事業の制限", 1), ("開発許可：許可が必要かどうか", 2), ("開発許可：手続き・建築制限", 1),
         ("建築基準法：単体規定・建築確認", 2), ("建築基準法：道路規制", 1), ("建築基準法：用途制限", 2),
         ("建築基準法：建蔽率・容積率", 2), ("建築基準法：高さ制限・日影規制・防火地域", 2),
         ("国土利用計画法（事後届出・事前届出）", 2), ("農地法（3条・4条・5条）", 2),
         ("宅地造成及び特定盛土等規制法", 2), ("土地区画整理法（仮換地・換地処分）", 2), ("その他の法令上の制限", 1)]
KENRI = [("制限行為能力者", 2), ("意思表示（詐欺・強迫・錯誤・虚偽表示）", 3), ("代理の基本", 2),
         ("無権代理・表見代理", 2), ("時効（取得時効・消滅時効）", 2), ("条件・期限", 1),
         ("債務不履行・契約の解除", 3), ("危険負担・手付", 1), ("売主の契約不適合責任", 2), ("弁済・相殺", 2),
         ("債権譲渡", 2), ("保証・連帯保証", 2), ("連帯債務", 1), ("物権変動と対抗要件（登記）", 3),
         ("不動産登記法", 3), ("抵当権の基本", 3), ("法定地上権・根抵当権", 2), ("留置権・先取特権・質権", 1),
         ("共有", 2), ("所有権・相隣関係", 1), ("区分所有法（マンション）", 3), ("賃貸借（民法）", 3),
         ("借地借家法：借家", 3), ("借地借家法：借地", 3), ("使用貸借", 1), ("請負・委任", 2),
         ("不法行為・使用者責任", 2), ("相続：相続人・相続分・承認と放棄", 3), ("遺言・遺留分・配偶者居住権", 2),
         ("事務管理・判決文問題の解き方", 1)]
ZEI = [("不動産取得税", 2), ("固定資産税", 2), ("所得税（譲渡所得の特例）", 3), ("印紙税", 2),
       ("登録免許税", 1), ("贈与税（住宅取得資金の特例）", 1), ("地価公示法", 2), ("不動産鑑定評価基準", 2),
       ("住宅金融支援機構", 2), ("景品表示法（不当表示）", 2), ("土地（地形・災害）", 2), ("建物（構造・材料）", 2),
       ("統計（今は出題形式を知るだけ）", 1)]

plan = {}   # 日付 -> 1日分


def put(d, phase, subj, theme, todo, memo="", score=False, gate=None, rest=False):
    assert d not in plan, f"{d} が重複"
    plan[d] = {"d": d.isoformat(), "phase": phase, "subj": subj, "theme": theme, "todo": todo,
               "memo": memo, "score": score, "gate": gate, "rest": rest}


def weekdays(start, end):
    """月〜土の勉強日(日曜と年末年始を除く)"""
    out, d = [], start
    while d <= end:
        if d.weekday() != 6 and d not in HOLIDAYS:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def allocate(topics, n):
    tot = sum(w for _, w in topics)
    raw = [w * n / tot for _, w in topics]
    cnt = [max(1, int(r)) for r in raw]
    order = sorted(range(len(topics)), key=lambda i: raw[i] - int(raw[i]), reverse=True)
    i = 0
    while sum(cnt) < n:
        cnt[order[i % len(order)]] += 1
        i += 1
    while sum(cnt) > n:
        cnt[max(range(len(cnt)), key=lambda k: cnt[k])] -= 1
    return cnt


def gate_day(d, phase, subj, theme, todo, items):
    put(d, phase, subj, "関門：" + theme, todo, score=False, gate=items,
        memo="届かなければ次に進まず、翌週はこの科目の復習にあてる")


def first_round(start, end, phase, subj, topics, gate_theme, gate_todo, gate_items, extra=None):
    days = weekdays(start, end)
    body, review, gday = days[:-2], days[-2], days[-1]
    cnt = allocate(topics, len(body))
    k = 0
    for (t, _), c in zip(topics, cnt):
        for p in range(c):
            if c == 1:
                main, th = f"②テキスト「{t}」を読み、分野別過去問の「{t}」を全部解く（2時間）", t
            elif p == 0:
                main, th = f"②テキスト「{t}」を最後まで読み、分野別過去問の「{t}」の前半を解く（2時間）", f"{t}（1/{c}日目）"
            elif p == c - 1:
                main, th = f"②分野別過去問の「{t}」の残りを解き、この単元の×をすべて解き直す（2時間）", f"{t}（{c}/{c}日目）"
            else:
                main, th = f"②分野別過去問の「{t}」の続きを解き、間違えた所をテキストで読み直す（2時間）", f"{t}（{p + 1}/{c}日目）"
            if extra:   # 4月: 税・その他と並行して宅建業法の2周目
                main = main.replace("（2時間）", "（1.5時間）")
                todo = f"{IPPON}\n{main}\n③宅建業法の2周目：分野別過去問の「{extra[k % len(extra)]}」を解き直す（1時間）"
            else:
                todo = f"{IPPON}\n{main}\n{FIX}"
            put(body[k], phase, subj, th, todo)
            k += 1
    put(review, phase, subj, f"{subj} 関門の前の総復習",
        f"{IPPON}\n②{subj}で×をつけた過去問をすべて解き直す（2時間）\n③まだ間違える問題を弱点ノートに書き写す（30分）")
    gate_day(gday, phase, subj, gate_theme, gate_todo, gate_items)


def g(gid, label, typ, target, frm=None):
    return {"id": gid, "label": label, "type": typ, "target": target, "from": frm}


# ===== 10月〜4月: 各科目の1周目 =====
first_round(D(2026, 10, 16), D(2026, 11, 30), "①宅建業法", "宅建業法", GYOHO, "宅建業法 正答率80%",
            f"{IPPON}\n②宅建業法の分野別過去問からランダムに40問、時間を計って解く（2時間）\n③正答率を出し、アプリの「今日」に入力（30分）",
            [g("g1-gyoho", "宅建業法の分野別過去問の正答率", "pct", 80)])
first_round(D(2026, 12, 1), D(2026, 12, 31), "②法令上の制限", "法令上の制限", HOREI, "法令上の制限70%・宅建業法80%",
            f"{IPPON}\n②法令上の制限と宅建業法の分野別過去問を、20問ずつランダムに解く（2時間）\n③それぞれの正答率を出し、アプリの「今日」に入力（30分）",
            [g("g2-horei", "法令上の制限の正答率", "pct", 70), g("g2-gyoho", "宅建業法の正答率（維持）", "pct", 80)])
first_round(D(2027, 1, 2), D(2027, 3, 31), "③権利関係", "権利関係", KENRI, "権利関係60%・宅建業法85%",
            f"{IPPON}\n②権利関係と宅建業法の分野別過去問を、20問ずつランダムに解く（2時間）\n③それぞれの正答率を出し、アプリの「今日」に入力（30分）",
            [g("g3-kenri", "権利関係の正答率", "pct", 60), g("g3-gyoho", "宅建業法の正答率（維持）", "pct", 85)])
first_round(D(2027, 4, 1), D(2027, 4, 30), "④税・その他＋業法2周目", "税・その他", ZEI, "全科目1周完了・実力測定",
            f"{IPPON}\n②年度別過去問（2026年度）50問を本番どおり2時間で解く\n③採点し、点数をアプリの「今日」に入力（30分）",
            [g("g4-first", "初めての年度別過去問（実力測定）", "score", None)], extra=[t for t, _ in GYOHO])

# ===== 5〜6月: 全科目2周目(平日) + 土曜に年度別の通し演習 =====
def chunk(topics, size):
    names = [t for t, _ in topics]
    return [names[i:i + size] for i in range(0, len(names), size)]


groups = ([("宅建業法", x) for x in chunk(GYOHO, 3)] + [("法令上の制限", x) for x in chunk(HOREI, 2)]
          + [("権利関係", x) for x in chunk(KENRI, 2)] + [("税・その他", x) for x in chunk(ZEI, 3)])
SAT_YEARS = ["2012年度", "2013年度", "2014年度", "2015年度", "2016年度", "2020年度（12月実施分）", "2021年度（12月実施分）"]
days = weekdays(D(2027, 5, 1), D(2027, 6, 30))
gday = days[-1]
gi = si = 0
for d in days[:-1]:
    if d.weekday() == 5:
        if si < len(SAT_YEARS):
            y = SAT_YEARS[si]
            put(d, "⑤全科目2周目", "過去問", f"通し演習：{y}",
                f"{IPPON}\n②{y}の過去問50問を本番どおり2時間で解く\n③採点して点数をアプリに入力し、×の解説を読む（30分）", score=True)
        else:
            put(d, "⑤全科目2周目", "過去問", "通し演習：苦手科目",
                f"{IPPON}\n②正答率がいちばん低い科目の分野別過去問を、時間を計って解く（2時間）\n{FIX}")
        si += 1
        continue
    subj, grp = groups[gi % len(groups)]
    rnd = "2周目" if gi < len(groups) else "3周目（×の問題だけ）"
    put(d, "⑤全科目2周目", subj, f"{rnd}：" + "／".join(grp),
        f"{IPPON}\n②分野別過去問の「{'」「'.join(grp)}」を解く（2時間）\n③×の問題をテキストで確認し、弱点ノートに1行メモ（30分）")
    gi += 1
gate_day(gday, "⑤全科目2周目", "過去問", "年度別過去問で33点以上",
         f"{IPPON}\n②2025年度の過去問50問を本番どおり2時間で解く\n③採点し、点数をアプリの「今日」に入力（30分）",
         [g("g5-33", "年度別過去問の点数", "score", 33)])

# ===== 7〜8月: 年度別過去問を2周 =====
YEARS = [f"{y}年度" for y in range(2017, 2027)]


def nendo(start, end, phase, label, gate):
    days = weekdays(start, end)
    body, gday = days[:-1], days[-1]
    for i, d in enumerate(body):
        memo = {D(2027, 7, 1): "⚠️ 今日中に試験の申込期間を確認する（不動産適正取引推進機構のサイト）",
                D(2027, 7, 12): "⚠️ 受験の申し込みは済みましたか？まだなら今日申し込む",
                D(2027, 7, 15): "⚠️ 8〜9月の会場模試（公開模試）を申し込む"}.get(d, "")
        if i < len(YEARS) * 2:
            y = YEARS[i // 2]
            if i % 2 == 0:
                put(d, phase, "過去問", f"{label}：{y}を解く",
                    f"{IPPON}\n②{y}の過去問50問を本番どおり2時間で解く\n③採点して、点数をアプリに入力（30分）", memo, score=True)
            else:
                put(d, phase, "過去問", f"{label}：{y}の復習",
                    f"{IPPON}\n②昨日の{y}の、全部の選択肢について「なぜ〇か×か」を説明できるまで解説を読む（2時間）\n{FIX}", memo)
        else:
            put(d, phase, "弱点", "弱点つぶし",
                f"{IPPON}\n②今月の過去問で点が低かった科目の分野別過去問を解き直す（2時間）\n③弱点ノートを読み返す（30分）", memo)
    gate_day(gday, phase, "過去問", gate[0], gate[1], gate[2])


nendo(D(2027, 7, 1), D(2027, 7, 31), "⑥年度別過去問1周目", "1周目",
      ("まだ届いていなくてよい", "", None))
plan.pop(max(d for d in plan if d.month == 7 and d.year == 2027))   # 7月は関門なし。最後の日は弱点つぶしにする
put(weekdays(D(2027, 7, 1), D(2027, 7, 31))[-1], "⑥年度別過去問1周目", "弱点", "弱点つぶし",
    f"{IPPON}\n②今月の過去問で点が低かった科目の分野別過去問を解き直す（2時間）\n③弱点ノートを読み返す（30分）")
nendo(D(2027, 8, 1), D(2027, 8, 31), "⑦年度別過去問2周目", "2周目",
      ("年度別過去問で37点以上を3回連続",
       f"{IPPON}\n②今月の点数を見直し、直近3回が37点以上か確かめる（30分）\n③届いていない科目の分野別過去問を解き直す（2時間）",
       [g("g7-37", "直近3回の年度別過去問（すべて37点以上）", "streak3", 37, "2027-08-01")]))

# ===== 9月: 模試・法改正・統計 =====
sep = weekdays(D(2027, 9, 1), D(2027, 9, 30))
items9 = []
for m in range(1, 5):
    items9 += [("模試", f"模試 第{m}回を解く", f"{IPPON}\n②市販の模試 第{m}回を本番どおり2時間で解く\n③採点して、点数をアプリに入力（30分）", True),
               ("模試", f"模試 第{m}回の復習", f"{IPPON}\n②模試 第{m}回で間違えた問題の解説を全部読む（1.5時間）\n③同じ単元の分野別過去問を解き直す（1時間）", False)]
items9 += [("法改正", "法改正のチェック", f"{IPPON}\n②テキスト・模試の「法改正」ページを読み、変わった点をノートにまとめる（2.5時間）", False),
           ("法改正", "法改正の問題演習", f"{IPPON}\n②法改正の予想問題を解く（2時間）\n③弱点ノートを読み返す（30分）", False),
           ("統計", "統計の暗記", f"{IPPON}\n②最新の統計（地価公示・着工戸数・取引件数など）を覚える（1時間）\n③宅建業法の×問題を解き直す（1.5時間）", False)]
weak = ["宅建業法", "法令上の制限", "権利関係", "税・その他"]
while len(items9) < len(sep) - 1:
    s = weak[(len(items9) - 11) % 4]
    items9.append(("弱点", f"弱点つぶし：{s}", f"{IPPON}\n②{s}の弱点ノートと×問題を解き直す（2時間）\n③統計を見直す（30分）", False))
order = [0, 1, 8, 2, 3, 11, 9, 4, 5, 12, 10, 6, 7] + list(range(13, len(items9)))
for d, idx in zip(sep[:-1], order):
    s, th, todo, sc = items9[idx]
    put(d, "⑧模試・法改正・統計", s, th, todo, score=sc)
gate_day(sep[-1], "⑧模試・法改正・統計", "模試", "模試で40点以上を3回連続",
         f"{IPPON}\n②9月の模試の点数を見直し、直近3回が40点以上か確かめる（30分）\n③届いていない科目の×問題を解き直す（2時間）",
         [g("g8-40", "直近3回の模試（すべて40点以上）", "streak3", 40, "2027-09-01")])

# ===== 10月: 直前期 =====
final = [("宅建業法", "総復習：宅建業法①", "業法の弱点ノートを読む→×問題を解き直す（3時間）"),
         ("宅建業法", "総復習：宅建業法②", "35条・37条・8種制限・報酬の×問題を解き直す（3時間）"),
         ("法令上の制限", "総復習：法令上の制限", "数字（面積・届出期限など）を一覧で覚え直す→×問題を解く（3時間）"),
         ("権利関係", "総復習：権利関係①", "よく出る単元（借地借家法・相続・抵当権・区分所有法）の×問題（3時間）"),
         ("権利関係", "総復習：権利関係②", "意思表示・代理・時効・物権変動の×問題（3時間）"),
         ("税・その他", "総復習：税・その他", "税の数字（税率・控除額）を覚え直す→統計の最終確認（3時間）"),
         ("過去問", "最新年度の過去問をもう一度", "2026年度の過去問を2時間で解く→復習（1時間）"),
         ("宅建業法", "総復習：宅建業法③", "業法の問題を1日で50問解く（3時間）"),
         ("全科目", "弱点ノート総ざらい①", "弱点ノートを最初から全部読む（3時間）"),
         ("全科目", "弱点ノート総ざらい②", "弱点ノートの中で、まだ覚えていない所だけ覚え直す（3時間）"),
         ("統計", "統計・法改正の最終暗記", "統計と法改正を声に出して覚える（3時間）"),
         ("全科目", "本番のシミュレーション", "本番と同じ13:00〜15:00に模試か過去問を1回解く→復習"),
         ("全科目", "軽い総復習", "弱点ノートを読むだけ（2時間）。新しい問題には手を出さない"),
         ("前日", "試験前日", "弱点ノートと統計を1時間だけ見る。受験票・筆記用具・腕時計を準備して早く寝る")]
for d, (s, th, todo) in zip(weekdays(D(2027, 10, 1), D(2027, 10, 16)), final):
    put(d, "⑨直前期", s, th, todo)
put(EXAM, "本番", "試験", "宅建試験 本番", "13:00〜15:00。宅建業法から解き始める", "ここまでやった自分を信じて")

# ===== 日曜: 週テスト(月の最初の日曜だけ完全休養)・年末年始 =====
d = START
while d < EXAM:
    if d not in plan:
        if d in HOLIDAYS:
            put(d, "", "休み", "年末年始の休み", "休み。余裕があれば一問一答だけ", rest=True)
        elif d.weekday() == 6 and d.day <= 7 and not (d.year == 2027 and d.month == 10):
            put(d, "", "休み", "完全休養日（月1回）", "何もしない日。しっかり休んで来週に備える", rest=True)
        elif d.weekday() == 6:
            put(d, "", "週テスト", "週テスト・今週の×直し",
                f"{IPPON}\n②週テスト：今週の範囲の分野別過去問を、時間を計って解く（1時間）\n③今週×をつけた問題をすべて解き直す（30分）")
    d += dt.timedelta(days=1)

# 最初の週は教材をそろえる
plan[START]["memo"] = "⚠️ 市販の分野別過去問集とテキスト（2027年度版。未発売なら過去問集は2026年度版でよい）を今週中に用意する"

rows = [plan[k] for k in sorted(plan)]
assert all((D.fromisoformat(r["d"]) - START).days == i for i, r in enumerate(rows)), "日付に抜けがある"
OUT_JSON.write_text(json.dumps(rows, ensure_ascii=False, indent=0), encoding="utf-8")

# ===== Excel(人が見る用) =====
wb = Workbook()
top = wb.active
top.title = "今日やること"
ws = wb.create_sheet("毎日の計画")
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
base_font = Font(name="メイリオ", size=10)
wrap = Alignment(wrap_text=True, vertical="center")
headers = ["日付", "曜日", "試験まで", "時期", "科目", "今日のテーマ", "やること（平日3時間・日曜2時間）", "完了", "メモ", "関門"]
widths = [12, 6, 9, 20, 12, 38, 64, 7, 30, 34]
for c, (h, w) in enumerate(zip(headers, widths), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill, cell.font = PatternFill("solid", fgColor="0F7B5F"), Font(bold=True, color="FFFFFF", name="メイリオ")
    cell.alignment, cell.border = Alignment(horizontal="center", vertical="center"), border
    ws.column_dimensions[cell.column_letter].width = w
for i, r in enumerate(rows, 2):
    d = D.fromisoformat(r["d"])
    gate = "\n".join(f"{x['label']}：" + (f"{x['target']}{'%' if x['type'] == 'pct' else '点'}以上" if x["target"] else "記録する")
                     for x in (r["gate"] or []))
    vals = [d, WD[d.weekday()], f"あと{(EXAM - d).days}日", r["phase"], r["subj"], r["theme"], r["todo"], "", r["memo"], gate]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font, cell.alignment, cell.border = base_font, wrap, border
    ws.cell(row=i, column=1).number_format = "yyyy/mm/dd"
    for c in (2, 3, 8):
        ws.cell(row=i, column=c).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[i].height = 15 * max(2, r["todo"].count("\n") + 1) + 4
last = len(rows) + 1
dv = DataValidation(type="list", formula1='"✓"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"H2:H{last}")
rng = f"A2:J{last}"
ws.conditional_formatting.add(rng, FormulaRule(formula=["$A2=TODAY()"], fill=PatternFill("solid", fgColor="FFE699"), font=Font(bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=['$H2="✓"'], fill=PatternFill("solid", fgColor="E2EFDA"), font=Font(color="808080")))
ws.conditional_formatting.add(rng, FormulaRule(formula=['$E2="休み"'], fill=PatternFill("solid", fgColor="EDEDED")))
ws.conditional_formatting.add(rng, FormulaRule(formula=['$J2<>""'], fill=PatternFill("solid", fgColor="FCE4D6"), font=Font(bold=True)))
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:J{last}"

top.column_dimensions["A"].width = 16
top.column_dimensions["B"].width = 80
R = f"毎日の計画!$A$2:$A${last}"


def lk(col):
    return f'=IFERROR(INDEX(毎日の計画!${col}$2:${col}${last},MATCH(TODAY(),{R},0)),"")'


top["A1"] = "宅建 2027年合格プラン（厳しめ版）"
top["A1"].font = Font(name="メイリオ", size=18, bold=True, color="0F7B5F")
items = [("今日の日付", "=TODAY()"), ("試験まで", f'="あと"&(DATE(2027,10,17)-TODAY())&"日"'),
         ("時期", lk("D")), ("科目", lk("E")), ("今日のテーマ", lk("F")), ("やること", lk("G")), ("メモ", lk("I")), ("関門", lk("J")),
         ("", ""), ("ふだんはアプリを使う", "https://ry1124.github.io/takken-study/ （このExcelは計画の一覧として残しています）")]
for r_, (k, v) in enumerate(items, 3):
    a, b = top.cell(row=r_, column=1, value=k), top.cell(row=r_, column=2, value=v)
    a.font = Font(name="メイリオ", size=11, bold=True, color="0F7B5F")
    b.font = Font(name="メイリオ", size=13 if k in ("今日のテーマ", "やること") else 11, bold=k == "今日のテーマ")
    a.alignment, b.alignment = Alignment(vertical="top"), Alignment(wrap_text=True, vertical="top")
top["B3"].number_format = "yyyy/mm/dd(aaa)"
top.row_dimensions[8].height = 70
os.makedirs(OUT_XLSX.parent, exist_ok=True)
wb.save(OUT_XLSX)

from collections import Counter
print(f"{len(rows)}日分 → {OUT_JSON.name} / {OUT_XLSX.name}")
print(Counter(r["subj"] for r in rows))
print("関門:", [(r["d"], r["theme"]) for r in rows if r["gate"]])
